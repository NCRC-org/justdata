"""Every lender and bank table states the year it covers, beside the table
(Jay round 2 #11): the DC LendSight report didn't make clear that its
top-lender table was 2025 only."""

import re
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from justdata.shared.pdf.components import latest_year

ROOT = Path(__file__).resolve().parents[2]
YEARS = [2021, 2022, 2023, 2024, 2025]


@pytest.mark.parametrize("years, expected", [
    (YEARS, "2025"), ([2025, 2021], "2025"), ("2021,2022,2025", "2025"),
    ("2024", "2024"), (2024, "2024"), ([], ""), ("", ""), (None, ""),
])
def test_latest_year(years, expected):
    assert latest_year(years) == expected


def test_web_lender_table_titles_carry_the_year():
    ls = (ROOT / "justdata/apps/lendsight/templates/partials/lendsight_report_template.html").read_text()
    assert re.search(r'Table 1: Top Lenders by Total Loans \(<span data-ls="last-year"></span>\)', ls)
    assert "in the most recent year." not in ls
    bs = (ROOT / "justdata/apps/bizsight/templates/partials/bizsight_report_template.html").read_text()
    assert bs.count('Top Lenders - Number of Loans (<span data-bs="latest-year"></span>)') == 1
    assert bs.count('Top Lenders - Amount of Loans (<span data-bs="latest-year"></span>)') == 1
    br = (ROOT / "justdata/apps/branchsight/templates/partials/branchsight_report_template.html").read_text()
    assert 'Top Banks by Branch Count (<span data-br="bank-year"></span>)' in br
    js = (ROOT / "justdata/shared/web/static/js/branchsight/br_report.js").read_text()
    assert "fill('bank-year', last)" in js


def test_lendsight_excel_lender_sheet_names_the_year(tmp_path):
    from justdata.apps.lendsight.report_builder.excel_export.writer import save_mortgage_excel_report
    lenders = pd.DataFrame([{"Lender Name": "Example Bank", "Lender Type": "Bank", "Total Loans": 10}])
    path = tmp_path / "ls.xlsx"
    save_mortgage_excel_report({"top_lenders_detailed": lenders}, str(path),
                               {"years": YEARS, "counties": ["Cook County, Illinois"]})
    assert "Section 3 - Top Lenders (2025)" in load_workbook(path).sheetnames


def test_bizsight_excel_lender_sheets_name_the_year(tmp_path):
    from justdata.apps.bizsight.excel_export import save_bizsight_excel_report
    path = tmp_path / "bs.xlsx"
    save_bizsight_excel_report({"top_lenders_table": [{"Lender Name": "Example Bank", "Num Total": 3,
                                                       "Amt Total (in $000s)": 120}]},
                               str(path), {"county_name": "Cook County, Illinois", "years": [2020, 2024]})
    names = load_workbook(path).sheetnames
    assert "Top Lenders - Number (2024)" in names and "Top Lenders - Amount (2024)" in names


def test_pdf_lender_headings_name_the_year():
    for app, needle in (("lendsight", "Section 3: Top Mortgage Lenders ({lender_year})"),
                        ("bizsight", "Section 4: Top Lenders ({lender_year})"),
                        ("branchsight", "in {bank_year}, the most recent year in this report")):
        src = (ROOT / f"justdata/apps/{app}/pdf_report.py").read_text()
        assert needle in src, app
