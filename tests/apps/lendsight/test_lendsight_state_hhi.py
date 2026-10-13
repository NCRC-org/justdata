"""Statewide HHI beside the county figure (Workflow 13277160986)."""

from pathlib import Path

import pandas as pd

from justdata.apps.lendsight.report_builder.sections.concentration import (
    create_market_concentration_table, hhi_comparison_rows, statewide_market_concentration)

ROOT = Path(__file__).resolve().parents[3]
STATE_ROWS = [  # two lenders, 75/25 in 2024: HHI 75^2 + 25^2 = 6,250
    {"year": 2024, "loan_purpose": "1", "lender_name": "A", "total_loan_amount": 300},
    {"year": 2024, "loan_purpose": "1", "lender_name": "B", "total_loan_amount": 100},
]


def test_statewide_uses_the_county_method():
    county_rows = [{"geoid5": "17031", "county_state": "Cook County, Illinois"}]
    calls = []

    def fetch(fips, years, purpose):
        calls.append((fips, years, purpose))
        return STATE_ROWS
    out = statewide_market_concentration(county_rows, [2024], ["purchase"], fetch=fetch)
    assert calls == [("17", [2024], ["purchase"])]
    assert out[0]["state"] == "Illinois" and out[0]["state_fips"] == "17"
    expected = create_market_concentration_table(pd.DataFrame(STATE_ROWS), [2024])
    assert out[0]["rows"] == [{str(k): v for k, v in r.items()} for r in expected]
    assert round(out[0]["rows"][0]["2024"]) == 6250


def test_failed_state_query_never_fails_the_report():
    def fetch(*a):
        raise RuntimeError("bq down")
    assert statewide_market_concentration([{"geoid5": "12057", "county_state": "Hillsborough County, Florida"}],
                                          [2024], None, fetch=fetch) == []


def test_comparison_rows_pair_county_and_state():
    county = [{"Loan Purpose": "All Loans", "2024": 304.87}, {"Loan Purpose": "Refinance", "2024": None}]
    states = [{"state": "Illinois", "rows": [{"Loan Purpose": "All Loans", "2024": 199.81},
                                             {"Loan Purpose": "Refinance", "2024": None}]}]
    rows = hhi_comparison_rows(county, states, [2024], "Cook County, Illinois")
    assert rows == [{"Market": "Cook County, Illinois", "Loan Purpose": "All Loans", "2024": 305},
                    {"Market": "Illinois (statewide)", "Loan Purpose": "All Loans", "2024": 200}]


def test_narrative_may_compare_but_not_explain():
    from justdata.apps.lendsight.analysis import LendSightAnalyzer
    analyzer = object.__new__(LendSightAnalyzer)
    block = analyzer._market_concentration_block([{"Loan Purpose": "All Loans", "2024": 305}],
                                                 [{"state": "Illinois", "rows": [{"Loan Purpose": "All Loans", "2024": 200}]}])
    assert "more or less concentrated than its state" in block
    assert "MUST NOT suggest why" in block
    assert "Statewide HHI" not in analyzer._market_concentration_block([], None)


def test_page_workbook_and_methods_show_the_state_figure(tmp_path):
    tpl = (ROOT / "justdata/apps/lendsight/templates/partials/lendsight_report_template.html").read_text()
    methods = (ROOT / "justdata/apps/lendsight/templates/partials/lendsight_methods.html").read_text()
    js = (ROOT / "justdata/shared/web/static/js/lendsight/ls_report.js").read_text()
    assert 'id="lsHhiCompare"' in tpl and "market_concentration_state" in js
    assert "statewide HHI applies the same calculation" in methods

    from openpyxl import load_workbook
    from justdata.apps.lendsight.report_builder.excel_export.writer import save_mortgage_excel_report
    path = tmp_path / "ls.xlsx"
    save_mortgage_excel_report(
        {"market_concentration": [{"Loan Purpose": "All Loans", "2024": 304.87}],
         "market_concentration_state": [{"state": "Illinois", "rows": [{"Loan Purpose": "All Loans", "2024": 199.81}]}]},
        str(path), {"years": [2024], "counties": ["Cook County, Illinois"]})
    ws = load_workbook(path)["Section 4 - Concentration (HHI)"]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    assert values[1] == ["Cook County, Illinois", "All Loans", 305]
    assert values[2] == ["Illinois (statewide)", "All Loans", 200]
