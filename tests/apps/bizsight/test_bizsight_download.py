"""BizSight /download: works from the stored result alone (no per-instance
progress record), names the county, and "excel" is the workbook alone."""

import io
import zipfile

import pytest

RESULT = {
    "metadata": {"counties": [{"name": "Warren County, Kentucky", "geoid5": "21227"}],
                 "county_name": "Warren County, Kentucky", "state_name": "Kentucky",
                 "years": [2020, 2021, 2022, 2023, 2024]},
    "county_summary_table": [], "summary_table": {},
}


@pytest.fixture
def client(monkeypatch):
    from justdata.main.app import create_app
    import justdata.apps.bizsight.blueprint as bp
    import justdata.apps.bizsight.excel_export as xl
    import justdata.apps.bizsight.pdf_report as pdf

    seen = {}
    monkeypatch.setattr(bp, "get_analysis_result_by_job_id", lambda job_id: RESULT)

    def fake_xlsx(result, path, metadata=None):
        seen["metadata"] = metadata
        open(path, "wb").write(b"PK\x03\x04xlsx")

    monkeypatch.setattr(xl, "save_bizsight_excel_report", fake_xlsx)
    monkeypatch.setattr(pdf, "generate_bizsight_pdf", lambda *a: io.BytesIO(b"%PDF-1.4"))
    app = create_app()
    app.config["TESTING"] = True
    c = app.test_client()
    with c.session_transaction() as s:   # no county_data / years in the session
        s["user_type"] = "staff"
        s["firebase_user"] = {"uid": "t", "email": "t@example.org", "email_verified": True}
    c.seen = seen
    return c


def test_excel_without_a_progress_record_names_the_county(client):
    resp = client.get("/bizsight/download?format=excel&job_id=never-tracked-here")
    assert resp.status_code == 200
    assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.data == b"PK\x03\x04xlsx"
    assert "Warren_County_Kentucky" in resp.headers["Content-Disposition"]
    assert client.seen["metadata"]["county_name"] == "Warren County, Kentucky"
    assert client.seen["metadata"]["state_name"] == "Kentucky"


def test_zip_bundles_workbook_and_pdf(client):
    resp = client.get("/bizsight/download?format=zip&job_id=x")
    names = zipfile.ZipFile(io.BytesIO(resp.data)).namelist()
    assert any(n.endswith(".xlsx") for n in names) and any(n.endswith(".pdf") for n in names)


def test_pdf(client):
    resp = client.get("/bizsight/download?format=pdf&job_id=x")
    assert resp.status_code == 200 and resp.data.startswith(b"%PDF")


def test_workbook_names_the_years_it_holds(tmp_path):
    """The workbook said 2018-2024 and "% Change Since 2018"; the change is
    measured from 2020 and the data runs 2020-2024."""
    import pandas as pd
    from openpyxl import load_workbook
    from justdata.apps.bizsight.excel_export import save_bizsight_excel_report
    from justdata.apps.bizsight.report_builder import create_comparison_table
    from tests.apps.bizsight.test_bizsight_denominator import ROW

    comparison = create_comparison_table(pd.DataFrame([ROW]), {}, {}, {'total_loans': 80}).to_dict('records')
    path = tmp_path / "bs.xlsx"
    save_bizsight_excel_report({'comparison_table': comparison}, str(path), RESULT["metadata"])
    wb = load_workbook(path)
    notes = " ".join(str(c.value) for row in wb["Methods & Definitions"].iter_rows() for c in row if c.value)
    assert "2018" not in notes and "aggregate data for 2020-2024" in notes and "HHI by year from 2020 to 2024" in notes
    header = [c.value for c in wb["Comparison - Number"][1]]
    assert "% Change Since 2020" in header and "% Change Since 2018" not in header
