"""BizSight statewide HHI beside the county figure (Workflow 13277160986)."""

from pathlib import Path

import pandas as pd

from justdata.apps.bizsight.report_builder import calculate_hhi_for_lenders, hhi_from_amounts, statewide_hhi_by_year

ROOT = Path(__file__).resolve().parents[3]


def test_state_and_county_use_one_calculation():
    county = pd.DataFrame([
        {"year": 2024, "respondent_id": "1", "lender_name": "A", "amt_under_100k": 300, "amt_100k_250k": 0, "amt_250k_1m": 0},
        {"year": 2024, "respondent_id": "2", "lender_name": "B", "amt_under_100k": 100, "amt_100k_250k": 0, "amt_250k_1m": 0},
    ])
    assert calculate_hhi_for_lenders(county, 2024) == hhi_from_amounts({"A": 300, "B": 100}) == 6250.0
    state = statewide_hhi_by_year([{"year": 2024, "lender_name": "A", "amt": 300},
                                   {"year": 2024, "lender_name": "B", "amt": 100}], [2023, 2024])
    assert state == [{"year": 2023, "hhi_value": None}, {"year": 2024, "hhi_value": 6250.0}]


def test_statewide_query_dedupes_like_the_county():
    from justdata.apps.bizsight.utils.bigquery_client import BigQueryClient
    import inspect
    sql = inspect.getsource(BigQueryClient.get_state_lender_amounts)
    assert "PARTITION BY d.year, d.respondent_id, LPAD(CAST(d.geoid5 AS STRING), 5, '0')" in sql
    assert "@state_fips" in sql and "UNNEST(@years)" in sql


def test_narrative_may_compare_but_not_explain():
    from justdata.apps.bizsight.ai_analysis import BizSightAnalyzer
    text = BizSightAnalyzer._statewide_hhi_text([{"year": 2024, "hhi_value": 517, "state_hhi_value": 377, "state_name": "Illinois"}])
    assert "more or less concentrated than Illinois" in text and "MUST NOT suggest why" in text
    assert BizSightAnalyzer._statewide_hhi_text([{"year": 2024, "hhi_value": 517}]) == ""


def test_page_workbook_and_methods_show_the_state_figure(tmp_path):
    tpl = (ROOT / "justdata/apps/bizsight/templates/partials/bizsight_report_template.html").read_text()
    js = (ROOT / "justdata/shared/web/static/js/bizsight/bs_report.js").read_text()
    methods = (ROOT / "justdata/apps/bizsight/templates/partials/bizsight_methods.html").read_text()
    assert 'id="bsHhiCompare"' in tpl and "state_hhi_value" in js
    assert "statewide HHI applies the same calculation" in methods

    from openpyxl import load_workbook
    from justdata.apps.bizsight.excel_export import save_bizsight_excel_report
    path = tmp_path / "bs.xlsx"
    save_bizsight_excel_report({"hhi_by_year": [{"year": 2024, "hhi_value": 517.37, "concentration_level": "Low",
                                                 "state_name": "Illinois", "state_hhi_value": 377.37}]},
                               str(path), {"county_name": "Cook County, Illinois", "years": [2024]})
    rows = list(load_workbook(path)["HHI by Year"].iter_rows(values_only=True))
    assert rows[0] == ("Year", "County HHI", "County concentration", "State", "Statewide HHI")
    assert rows[1] == (2024, 517.37, "Low", "Illinois", 377.37)
